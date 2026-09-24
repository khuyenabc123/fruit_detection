from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("YOLOv8_Tim_hieu_va_Ung_dung.pptx")

W = 13.333
H = 7.5

# Light presentation theme. WHITE is intentionally the primary ink color name
# retained from the original deck so all existing text calls stay concise.
NAVY = RGBColor(247, 249, 252)
NAVY_2 = RGBColor(238, 243, 249)
CARD = RGBColor(255, 255, 255)
CARD_2 = RGBColor(245, 248, 252)
CYAN = RGBColor(0, 132, 145)
GREEN = RGBColor(18, 137, 76)
YELLOW = RGBColor(190, 112, 0)
RED = RGBColor(197, 53, 70)
PURPLE = RGBColor(100, 75, 196)
WHITE = RGBColor(22, 38, 61)
MUTED = RGBColor(82, 100, 124)
GRID = RGBColor(214, 223, 233)
CODE_BG = RGBColor(20, 31, 48)
LIGHT_TEXT = RGBColor(245, 248, 252)

FONT = "Noto Sans"
MONO = "DejaVu Sans Mono"


def rect(slide, x, y, w, h, fill, radius=False, line=None, transparency=0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.fill.transparency = transparency
    shape.line.color.rgb = line or fill
    return shape


def line(slide, x1, y1, x2, y2, color=GRID, width=1.5):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x1), Inches(y1), Inches(max(0.01, x2 - x1)), Inches(max(0.01, y2 - y1))
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.color.rgb = color
    return shp


def text(slide, value, x, y, w, h, size=20, color=WHITE, bold=False,
         align=PP_ALIGN.LEFT, font=FONT, valign=MSO_ANCHOR.TOP,
         margin=0.04, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    frame.word_wrap = True
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def rich_text(slide, parts, x, y, w, h, size=20, color=WHITE,
              align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    frame.vertical_anchor = valign
    frame.word_wrap = True
    p = frame.paragraphs[0]
    p.alignment = align
    for value, part_color, bold in parts:
        run = p.add_run()
        run.text = value
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = part_color or color
    return box


def bullet_list(slide, items, x, y, w, h, size=18, color=WHITE,
                bullet_color=CYAN, spacing=8):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = 0
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = Pt(spacing)
        p.line_spacing = 1.08
        r1 = p.add_run()
        r1.text = "●  "
        r1.font.name = FONT
        r1.font.size = Pt(max(9, size - 5))
        r1.font.color.rgb = bullet_color
        r2 = p.add_run()
        r2.text = item
        r2.font.name = FONT
        r2.font.size = Pt(size)
        r2.font.color.rgb = color
    return box


def title(slide, kicker, heading, number):
    text(slide, kicker.upper(), 0.72, 0.38, 6.8, 0.3, 10, CYAN, True)
    text(slide, heading, 0.72, 0.72, 11.8, 0.72, 26, WHITE, True)
    text(slide, f"{number:02d}", 12.06, 0.50, 0.55, 0.35, 11, MUTED, True, PP_ALIGN.RIGHT)
    line(slide, 0.72, 1.46, 12.6, 1.475, GRID)


def footer(slide, number, source=None):
    if source:
        text(slide, source, 0.72, 7.12, 10.9, 0.2, 7.5, MUTED)
    text(slide, f"YOLOv8  ·  {number:02d}", 11.55, 7.10, 1.05, 0.22, 8, MUTED, True, PP_ALIGN.RIGHT)


def add_bg(slide):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY
    # A subtle corner accent.
    rect(slide, 11.9, -0.28, 1.8, 1.05, CYAN, radius=True, transparency=82)
    rect(slide, -0.35, 6.85, 2.1, 0.9, PURPLE, radius=True, transparency=88)


def card(slide, x, y, w, h, heading, body, accent=CYAN, number=None):
    rect(slide, x, y, w, h, CARD, radius=True, line=GRID)
    rect(slide, x, y, 0.055, h, accent, radius=True, line=accent)
    if number:
        text(slide, number, x + 0.22, y + 0.18, 0.4, 0.38, 12, accent, True)
        hx = x + 0.68
        hw = w - 0.9
    else:
        hx = x + 0.28
        hw = w - 0.55
    text(slide, heading, hx, y + 0.17, hw, 0.38, 15, WHITE, True)
    text(slide, body, x + 0.28, y + 0.72, w - 0.55, h - 0.9, 11.5, MUTED)


def pill(slide, label, x, y, w, color=CYAN):
    rect(slide, x, y, w, 0.38, color, radius=True, line=color, transparency=78)
    text(slide, label, x + 0.07, y + 0.065, w - 0.14, 0.23, 9.2, WHITE, True, PP_ALIGN.CENTER)


def add_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    # Abstract feature-map grid.
    for i in range(6):
        for j in range(5):
            c = CYAN if (i + j) % 4 == 0 else GRID
            rect(slide, 8.32 + i * 0.64, 1.26 + j * 0.64, 0.48, 0.48, c, radius=True,
                 transparency=25 if c == CYAN else 10)
    # Bounding-box motif.
    outline = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.05), Inches(2.03), Inches(2.6), Inches(2.12))
    outline.fill.background()
    outline.line.color.rgb = GREEN
    outline.line.width = Pt(3)
    pill(slide, "APPLE  0.94", 9.05, 1.66, 1.52, GREEN)
    text(slide, "MACHINE LEARNING · COMPUTER VISION", 0.86, 0.74, 6.8, 0.34, 11, CYAN, True)
    text(slide, "TÌM HIỂU\nYOLOv8", 0.86, 1.52, 7.1, 2.0, 39, WHITE, True)
    text(slide, "Kiến trúc · Huấn luyện · Đánh giá · Triển khai", 0.90, 3.73, 6.8, 0.5, 18, MUTED)
    line(slide, 0.90, 4.42, 6.7, 4.435, GRID)
    text(slide, "Ứng dụng minh họa: Fruit Detection", 0.90, 4.68, 5.6, 0.4, 15, GREEN, True)
    text(slide, "Group Project  ·  2026", 0.90, 6.63, 4.0, 0.3, 10, MUTED, True)
    return slide


def add_detection_overview(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    title(slide, "01 · Bối cảnh", "YOLO giải quyết bài toán gì?", 2)

    # Left: mock image with fruit detections.
    rect(slide, 0.72, 1.84, 6.05, 4.72, CARD, radius=True, line=GRID)
    rect(slide, 0.96, 2.10, 5.57, 4.10, RGBColor(225, 240, 230), radius=True, line=RGBColor(225, 240, 230))
    fruit = [
        (1.55, 2.80, 1.28, 1.28, RED, "APPLE 0.94"),
        (3.50, 3.58, 1.05, 1.05, YELLOW, "ORANGE 0.91"),
        (4.80, 2.52, 0.92, 0.92, GREEN, "LIME 0.87"),
    ]
    for fx, fy, fw, fh, fc, label in fruit:
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(fx), Inches(fy), Inches(fw), Inches(fh))
        circle.fill.solid(); circle.fill.fore_color.rgb = fc; circle.line.color.rgb = fc
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(fx - 0.12), Inches(fy - 0.25), Inches(fw + 0.24), Inches(fh + 0.38))
        box.fill.background(); box.line.color.rgb = CYAN; box.line.width = Pt(2)
        pill(slide, label, fx - 0.12, fy - 0.54, max(1.25, fw + 0.55), CYAN)
    text(slide, "ẢNH ĐẦU VÀO", 1.02, 6.24, 2.0, 0.22, 9, MUTED, True)

    rich_text(slide, [("Một lần suy luận", CYAN, True), (" → nhiều dự đoán", WHITE, False)],
              7.27, 1.98, 5.18, 0.48, 19)
    card(slide, 7.27, 2.63, 5.18, 0.96, "Vị trí", "Bounding box: x, y, width, height", CYAN, "01")
    card(slide, 7.27, 3.79, 5.18, 0.96, "Phân loại", "Tên lớp: apple, orange, banana…", GREEN, "02")
    card(slide, 7.27, 4.95, 5.18, 0.96, "Độ tin cậy", "Confidence score cho từng dự đoán", YELLOW, "03")
    footer(slide, 2, "Nguồn: Ultralytics YOLOv8 documentation")


def add_yolov8_intro(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    title(slide, "02 · Tổng quan", "YOLOv8: một nền tảng thị giác máy tính thống nhất", 3)
    text(slide, "Ra mắt", 0.78, 1.86, 1.4, 0.3, 10, MUTED, True)
    text(slide, "01 / 2023", 0.78, 2.17, 2.2, 0.55, 27, CYAN, True)
    text(slide, "Phát triển bởi Ultralytics", 0.78, 2.77, 2.5, 0.36, 12, WHITE, True)
    line(slide, 3.34, 1.86, 3.355, 6.42, GRID)

    card(slide, 3.76, 1.88, 4.12, 1.63, "Anchor-free split head",
         "Dự đoán trực tiếp tại các điểm đặc trưng; tách nhánh hồi quy box và phân loại.", CYAN)
    card(slide, 8.15, 1.88, 4.12, 1.63, "Accuracy ↔ Speed",
         "Nhiều kích cỡ model để chọn theo GPU, latency và yêu cầu độ chính xác.", GREEN)
    card(slide, 3.76, 3.78, 4.12, 1.63, "Một API, nhiều tác vụ",
         "Detect · Segment · Classify · Pose · OBB", PURPLE)
    card(slide, 8.15, 3.78, 4.12, 1.63, "Pipeline end-to-end",
         "Train · Validate · Predict · Export trong cùng hệ sinh thái Ultralytics.", YELLOW)
    text(slide, "Điểm cần nhớ", 3.78, 5.83, 1.6, 0.28, 11, CYAN, True)
    text(slide, "YOLOv8 không chỉ là một model detection — đây là một họ model và bộ công cụ triển khai.",
         5.35, 5.77, 6.95, 0.52, 14, WHITE, True)
    footer(slide, 3, "Nguồn: docs.ultralytics.com/models/yolov8")


def add_architecture(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    title(slide, "03 · Kiến trúc", "Backbone → Neck → Head", 4)

    stages = [
        (0.80, 2.04, 3.38, 3.88, "BACKBONE", "Trích xuất đặc trưng", CYAN,
         ["Conv + C2f", "SPPF", "Feature maps đa tỉ lệ"]),
        (4.95, 2.04, 3.38, 3.88, "NECK", "Hợp nhất thông tin", GREEN,
         ["FPN: top-down", "PAN: bottom-up", "P3 · P4 · P5"]),
        (9.10, 2.04, 3.38, 3.88, "HEAD", "Sinh dự đoán", PURPLE,
         ["Anchor-free", "Box / Class tách nhánh", "DFL + IoU-based loss"]),
    ]
    for x, y, w, h, name, desc, accent, items in stages:
        rect(slide, x, y, w, h, CARD, radius=True, line=GRID)
        pill(slide, name, x + 0.28, y + 0.28, 1.33, accent)
        text(slide, desc, x + 0.28, y + 0.91, w - 0.56, 0.4, 17, WHITE, True)
        bullet_list(slide, items, x + 0.30, y + 1.57, w - 0.60, 1.86, 13, MUTED, accent, 10)
        text(slide, "FEATURE FLOW", x + 0.30, y + 3.42, w - 0.60, 0.2, 8, accent, True)
    # Flow arrows.
    for ax in (4.33, 8.48):
        arrow = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(ax), Inches(3.48), Inches(0.42), Inches(0.70))
        arrow.fill.solid(); arrow.fill.fore_color.rgb = CYAN; arrow.fill.transparency = 15
        arrow.line.color.rgb = CYAN
    text(slide, "C2f giúp duy trì dòng gradient và tái sử dụng đặc trưng hiệu quả hơn block C3 của YOLOv5.",
         1.02, 6.27, 11.35, 0.47, 13, WHITE, True, PP_ALIGN.CENTER)
    footer(slide, 4, "Nguồn: Ultralytics YOLO Architecture Guide")


def add_innovations(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    title(slide, "04 · Cải tiến", "YOLOv8 thay đổi gì so với YOLOv5?", 5)

    headers = [("THÀNH PHẦN", 0.80, 2.12), ("YOLOv5", 4.12, 2.12), ("YOLOv8", 7.31, 2.12), ("Ý NGHĨA", 10.19, 2.12)]
    for label, x, y in headers:
        text(slide, label, x, y, 2.25, 0.30, 10, CYAN, True)
    rows = [
        ("Backbone block", "C3", "C2f", "Dòng gradient phong phú"),
        ("Detection head", "Anchor-based", "Anchor-free", "Bớt bước cấu hình anchor"),
        ("Nhánh dự đoán", "Gắn kết", "Tách box / class", "Chuyên biệt hóa nhiệm vụ"),
        ("Box regression", "IoU-based", "DFL + IoU-based", "Định vị box chi tiết hơn"),
        ("Workflow", "Repo riêng", "API thống nhất", "Train → export thuận tiện"),
    ]
    y = 2.64
    for i, row in enumerate(rows):
        fill = CARD if i % 2 == 0 else NAVY_2
        rect(slide, 0.76, y - 0.10, 11.82, 0.72, fill, radius=True, line=fill)
        xs = [0.95, 4.12, 7.31, 10.19]
        widths = [2.72, 2.55, 2.35, 2.16]
        colors = [WHITE, MUTED, GREEN, MUTED]
        bolds = [True, False, True, False]
        for val, x, w, c, b in zip(row, xs, widths, colors, bolds):
            text(slide, val, x, y + 0.04, w, 0.38, 12, c, b, valign=MSO_ANCHOR.MIDDLE)
        y += 0.84
    rect(slide, 0.90, 6.70, 11.52, 0.30, CYAN, radius=True, transparency=75)
    text(slide, "Kết quả: pipeline đơn giản hơn, linh hoạt hơn và dễ chọn điểm cân bằng accuracy–latency.",
         1.08, 6.72, 11.10, 0.23, 10.5, WHITE, True, PP_ALIGN.CENTER)
    footer(slide, 5, "Nguồn: Ultralytics YOLOv8 vs YOLOv5 comparison")


def add_variants(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    title(slide, "05 · Lựa chọn model", "Một kiến trúc, nhiều quy mô và tác vụ", 6)

    # Scale spectrum.
    text(slide, "QUY MÔ", 0.80, 1.86, 1.3, 0.26, 10, CYAN, True)
    scales = [
        ("n", "Nano", 0.78, 1.42, CYAN),
        ("s", "Small", 2.45, 1.62, GREEN),
        ("m", "Medium", 4.32, 1.82, YELLOW),
        ("l", "Large", 6.40, 2.04, RED),
        ("x", "X-Large", 8.70, 2.30, PURPLE),
    ]
    y = 2.38
    for letter, label, x, w, c in scales:
        rect(slide, x, y, w, 1.26, c, radius=True, line=c, transparency=72)
        text(slide, letter, x, y + 0.16, w, 0.44, 23, WHITE, True, PP_ALIGN.CENTER)
        text(slide, label, x, y + 0.72, w, 0.26, 9, MUTED, True, PP_ALIGN.CENTER)
    rich_text(slide, [("nhanh / nhẹ", CYAN, True), ("   →   ", MUTED, False), ("chính xác / nặng", PURPLE, True)],
              0.82, 3.90, 10.15, 0.34, 12, align=PP_ALIGN.CENTER)

    text(slide, "TÁC VỤ HỖ TRỢ", 0.80, 4.60, 2.2, 0.26, 10, CYAN, True)
    tasks = [("Detect", CYAN), ("Segment", GREEN), ("Classify", YELLOW), ("Pose", PURPLE), ("OBB", RED)]
    tx = 0.80
    for label, c in tasks:
        pill(slide, label, tx, 5.05, 1.62, c)
        tx += 1.83
    card(slide, 10.30, 4.58, 2.12, 1.34, "Gợi ý bắt đầu", "yolov8n để kiểm tra pipeline; yolov8s cho baseline tốt.", GREEN)
    text(slide, "Chọn model theo thiết bị triển khai và latency thực tế — không chỉ theo mAP trên COCO.",
         0.84, 6.34, 9.10, 0.42, 13, WHITE, True)
    footer(slide, 6, "Tên model ví dụ: yolov8n.pt · yolov8s-seg.pt · yolov8n-pose.pt")


def add_dataset(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    title(slide, "06 · Dữ liệu", "Dataset quyết định trần hiệu năng", 7)

    rect(slide, 0.76, 1.86, 5.24, 4.86, CARD, radius=True, line=GRID)
    text(slide, "CẤU TRÚC", 1.03, 2.12, 1.5, 0.28, 10, CYAN, True)
    structure = "dataset/\n├── images/\n│   ├── train/\n│   └── val/\n├── labels/\n│   ├── train/\n│   └── val/\n└── data.yaml"
    text(slide, structure, 1.04, 2.55, 2.30, 2.92, 15, WHITE, False, font=MONO)
    rect(slide, 3.46, 2.62, 2.15, 2.14, NAVY_2, radius=True, line=GRID)
    text(slide, "LABEL", 3.72, 2.88, 1.55, 0.25, 10, GREEN, True)
    text(slide, "class_id\nx_center  y_center\nwidth  height", 3.72, 3.35, 1.58, 1.0, 11, WHITE, False, font=MONO)
    text(slide, "Tọa độ chuẩn hóa 0–1", 3.70, 4.48, 1.72, 0.22, 8.5, MUTED)
    text(slide, "Ví dụ", 1.04, 5.76, 0.72, 0.24, 9, MUTED, True)
    text(slide, "0  0.52  0.41  0.20  0.16", 1.74, 5.73, 3.54, 0.30, 11, YELLOW, True, font=MONO)

    text(slide, "CHECKLIST CHO FRUIT DETECTION", 6.45, 1.98, 4.5, 0.28, 11, CYAN, True)
    checks = [
        "Đủ ánh sáng, góc chụp và khoảng cách",
        "Có quả nhỏ, bị che và chồng lấp",
        "Nhãn box nhất quán giữa người gán nhãn",
        "Không chia các frame gần nhau sang train/val",
        "Có ảnh nền không chứa trái cây",
        "Kiểm tra cân bằng giữa các lớp",
    ]
    bullet_list(slide, checks, 6.47, 2.50, 5.78, 3.80, 14, WHITE, GREEN, 11)
    rect(slide, 6.45, 6.24, 5.75, 0.42, RED, radius=True, transparency=76)
    text(slide, "Nhãn sai thường gây hại hơn việc chọn sai vài hyperparameter.", 6.64, 6.32, 5.37, 0.22, 9.8, WHITE, True, PP_ALIGN.CENTER)
    footer(slide, 7, "Nguồn: docs.ultralytics.com/datasets/detect")


def add_training(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    title(slide, "07 · Huấn luyện", "Fine-tune YOLOv8 với dữ liệu riêng", 8)

    steps = [
        ("01", "Chuẩn bị", "Ảnh + nhãn + data.yaml", CYAN),
        ("02", "Train", "Transfer learning từ pretrained weights", GREEN),
        ("03", "Validate", "Đọc P · R · mAP · confusion matrix", YELLOW),
        ("04", "Error analysis", "Xem false positive / false negative", PURPLE),
        ("05", "Export", "ONNX · TensorRT · OpenVINO", RED),
    ]
    sx = 0.75
    for num, head, body, c in steps:
        rect(slide, sx, 1.89, 2.28, 1.44, CARD, radius=True, line=GRID)
        text(slide, num, sx + 0.18, 2.06, 0.40, 0.28, 10, c, True)
        text(slide, head, sx + 0.62, 2.00, 1.40, 0.36, 13, WHITE, True)
        text(slide, body, sx + 0.18, 2.53, 1.89, 0.56, 9.5, MUTED)
        sx += 2.48

    rect(slide, 0.75, 3.68, 11.82, 2.45, CODE_BG, radius=True, line=GRID)
    text(slide, "CLI", 1.02, 3.95, 0.52, 0.25, 10, CYAN, True, font=MONO)
    code = (
        "$ pip install ultralytics\n"
        "$ yolo detect train model=yolov8s.pt data=data.yaml epochs=100 imgsz=640 batch=16\n"
        "$ yolo detect val model=runs/detect/train/weights/best.pt data=data.yaml\n"
        "$ yolo detect predict model=best.pt source=test.jpg conf=0.25"
    )
    text(slide, code, 1.02, 4.36, 10.98, 1.40, 12.2, LIGHT_TEXT, False, font=MONO)
    text(slide, "Bắt đầu bằng baseline mặc định → phân tích lỗi → chỉ thay đổi một nhóm yếu tố mỗi lần.",
         1.00, 6.40, 11.34, 0.35, 12, GREEN, True, PP_ALIGN.CENTER)
    footer(slide, 8, "Pretrained weights thường hội tụ nhanh hơn train from scratch")


def add_metrics(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    title(slide, "08 · Đánh giá", "Một con số mAP là chưa đủ", 9)

    metrics = [
        ("PRECISION", "TP / (TP + FP)", "Ít báo nhầm", CYAN),
        ("RECALL", "TP / (TP + FN)", "Ít bỏ sót", GREEN),
        ("IoU", "overlap / union", "Độ khớp box", YELLOW),
        ("mAP50–95", "mean AP @ IoU", "Đánh giá tổng hợp", PURPLE),
    ]
    mx = 0.76
    for name, formula, meaning, c in metrics:
        rect(slide, mx, 1.91, 2.83, 1.62, CARD, radius=True, line=GRID)
        text(slide, name, mx + 0.22, 2.12, 2.38, 0.28, 10, c, True)
        text(slide, formula, mx + 0.22, 2.56, 2.38, 0.34, 14.5, WHITE, True, font=MONO)
        text(slide, meaning, mx + 0.22, 3.03, 2.38, 0.25, 9.5, MUTED)
        mx += 3.02

    rect(slide, 0.76, 3.91, 7.17, 2.18, CARD, radius=True, line=GRID)
    text(slide, "Đọc lỗi theo mục tiêu", 1.03, 4.17, 3.0, 0.34, 16, WHITE, True)
    bullet_list(slide, [
        "Đếm quả: ưu tiên recall và sai số đếm",
        "Robot hái quả: ưu tiên IoU / mask chính xác",
        "Camera real-time: đo latency và FPS trên thiết bị thật",
    ], 1.02, 4.74, 6.40, 1.12, 12, MUTED, CYAN, 7)

    rect(slide, 8.27, 3.91, 4.30, 2.18, NAVY_2, radius=True, line=GRID)
    text(slide, "Ngưỡng confidence", 8.56, 4.18, 3.65, 0.34, 15, WHITE, True)
    rich_text(slide, [("Tăng", CYAN, True), (" conf → Precision ↑, Recall có thể ↓", MUTED, False)], 8.56, 4.79, 3.55, 0.42, 11.5)
    rich_text(slide, [("Giảm", GREEN, True), (" conf → Recall ↑, FP có thể ↑", MUTED, False)], 8.56, 5.30, 3.55, 0.42, 11.5)
    text(slide, "Chọn threshold trên validation set theo KPI nghiệp vụ.", 0.82, 6.42, 11.60, 0.35, 12, WHITE, True, PP_ALIGN.CENTER)
    footer(slide, 9, "Nguồn: Ultralytics YOLO Performance Metrics Guide")


def add_deployment(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    title(slide, "09 · Production", "Model tốt ≠ hệ thống production tốt", 10)

    stages = [
        ("Camera", "Input", CYAN),
        ("Preprocess", "Resize · color", GREEN),
        ("YOLOv8", "Inference", YELLOW),
        ("Postprocess", "NMS · threshold", PURPLE),
        ("Tracking", "Count · ID", RED),
        ("API / UI", "Output", CYAN),
    ]
    x = 0.61
    for i, (head, sub, c) in enumerate(stages):
        rect(slide, x, 2.03, 1.76, 1.16, CARD, radius=True, line=GRID)
        text(slide, head, x + 0.12, 2.26, 1.52, 0.28, 11.5, WHITE, True, PP_ALIGN.CENTER)
        text(slide, sub, x + 0.12, 2.72, 1.52, 0.20, 8.5, c, True, PP_ALIGN.CENTER)
        if i < len(stages) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x + 1.79), Inches(2.39), Inches(0.32), Inches(0.45))
            arrow.fill.solid(); arrow.fill.fore_color.rgb = GRID; arrow.line.color.rgb = GRID
        x += 2.12

    checklist = [
        ("Accuracy", "Test trên dữ liệu camera thật", CYAN),
        ("Latency", "Đo p95 / FPS trên phần cứng thật", GREEN),
        ("Reliability", "Timeout · health check · fallback", YELLOW),
        ("MLOps", "Version model · logging · rollback", PURPLE),
        ("Monitoring", "Data drift · confidence · lỗi", RED),
        ("Security", "Giới hạn file · quyền truy cập", CYAN),
    ]
    cx, cy = 0.76, 3.84
    for i, (head, body, c) in enumerate(checklist):
        card(slide, cx, cy, 3.74, 0.92, head, body, c)
        cx += 3.99
        if (i + 1) % 3 == 0:
            cx = 0.76
            cy += 1.12
    text(slide, "Nếu đếm trong video: detection cần kết hợp tracking để tránh đếm lặp cùng một quả.",
         0.86, 6.42, 11.55, 0.34, 12.5, GREEN, True, PP_ALIGN.CENTER)
    footer(slide, 10, "Định dạng triển khai phổ biến: PyTorch · ONNX · TensorRT · OpenVINO")


def add_summary(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    title(slide, "10 · Kết luận", "Ba thông điệp cần nhớ", 11)

    cards = [
        ("01", "Hiểu pipeline", "Nắm backbone–neck–head, preprocessing và postprocessing.", CYAN),
        ("02", "Ưu tiên dữ liệu", "Nhãn tốt, dữ liệu đa dạng và split đúng quan trọng hơn chạy theo version.", GREEN),
        ("03", "Đo theo nghiệp vụ", "mAP cần đi cùng recall, latency, FPS và sai số đếm thực tế.", PURPLE),
    ]
    y = 1.95
    for num, head, body, c in cards:
        rect(slide, 0.84, y, 8.02, 1.14, CARD, radius=True, line=GRID)
        text(slide, num, 1.10, y + 0.27, 0.58, 0.40, 15, c, True)
        text(slide, head, 1.86, y + 0.19, 2.24, 0.34, 15, WHITE, True)
        text(slide, body, 4.13, y + 0.20, 4.34, 0.58, 11.5, MUTED)
        y += 1.37

    rect(slide, 9.24, 1.95, 3.18, 3.88, NAVY_2, radius=True, line=GRID)
    text(slide, "NEXT STEP", 9.58, 2.28, 2.45, 0.28, 10, CYAN, True)
    text(slide, "Xây dựng baseline\nYOLOv8n / YOLOv8s\ntrên fruit dataset", 9.58, 2.89, 2.45, 1.60, 18, WHITE, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    pill(slide, "TRAIN → EVALUATE → ITERATE", 9.52, 5.00, 2.62, GREEN)

    text(slide, "TÀI LIỆU THAM KHẢO", 0.86, 6.18, 2.10, 0.23, 9, CYAN, True)
    text(slide, "docs.ultralytics.com/models/yolov8  ·  docs.ultralytics.com/guides/yolo-architecture  ·  docs.ultralytics.com/datasets/detect",
         0.86, 6.52, 11.58, 0.32, 8.5, MUTED)
    text(slide, "Q&A", 10.66, 6.49, 1.70, 0.42, 18, WHITE, True, PP_ALIGN.RIGHT)
    footer(slide, 11)


def set_properties(prs):
    prs.core_properties.title = "Tìm hiểu YOLOv8"
    prs.core_properties.subject = "Kiến trúc, huấn luyện, đánh giá và triển khai YOLOv8"
    prs.core_properties.author = "Machine Learning Group Project"
    prs.core_properties.keywords = "YOLOv8, object detection, fruit detection, machine learning"
    prs.core_properties.comments = "Generated as an editable PowerPoint deck."


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    set_properties(prs)

    add_cover(prs)
    add_detection_overview(prs)
    add_yolov8_intro(prs)
    add_architecture(prs)
    add_innovations(prs)
    add_variants(prs)
    add_dataset(prs)
    add_training(prs)
    add_metrics(prs)
    add_deployment(prs)
    add_summary(prs)

    prs.save(OUT)
    print(f"Created {OUT} with {len(prs.slides)} slides")


if __name__ == "__main__":
    build()
